"""產生「全班對戰」題庫的 Excel 範本：templates/quiz-template.xlsx

用法：python3 tools/make_quiz_template.py
"""
import pathlib
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = ROOT / 'templates' / 'quiz-template.xlsx'
ROWS = 100   # 範本預先設好格式的題數（題庫一個題組最多 100 題）

ROSE, ROSE_LIGHT, INK = 'C25D7C', 'FDE9EF', '3F2830'
thin = Side(style='thin', color='F0D6DE')

wb = Workbook()
ws = wb.active
ws.title = '題目'
headers = ['題目', '選項A', '選項B', '選項C（可空白）', '選項D（可空白）', '正確答案（A～D）']
widths = [46, 18, 18, 18, 18, 18]
for i, (h, w) in enumerate(zip(headers, widths), start=1):
    c = ws.cell(row=1, column=i, value=h)
    c.font = Font(bold=True, color='FFFFFF', size=12)
    c.fill = PatternFill('solid', fgColor=ROSE)
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws.column_dimensions[c.column_letter].width = w
ws.row_dimensions[1].height = 26
ws.freeze_panes = 'A2'

examples = [
    ['3/4 和哪一個分數一樣大？', '6/8', '4/5', '3/8', '6/4', 'A'],
    ['「閱讀」的「閱」是什麼部首？', '門', '兌', '目', '', 'A'],
    ['1/2 比 1/3 大', '⭕ 對', '❌ 錯', '', '', 'A'],
]
for r in range(2, ROWS + 2):
    for col in range(1, 7):
        c = ws.cell(row=r, column=col)
        c.number_format = '@'          # 文字格式：打 3/4 不會被 Excel 變成日期
        c.alignment = Alignment(vertical='center', wrap_text=(col == 1), horizontal='left' if col == 1 else 'center')
        c.border = Border(bottom=thin)
        if r % 2 == 1:
            c.fill = PatternFill('solid', fgColor=ROSE_LIGHT)
for r, row in enumerate(examples, start=2):
    for col, v in enumerate(row, start=1):
        ws.cell(row=r, column=col, value=v)

dv = DataValidation(type='list', formula1='"A,B,C,D"', allow_blank=True,
                    errorStyle='warning', showErrorMessage=True,
                    errorTitle='正確答案', error='正確答案請填 A、B、C 或 D（也可以直接寫正確選項的文字）。',
                    promptTitle='正確答案', prompt='填 A、B、C 或 D', showInputMessage=True)
ws.add_data_validation(dv)
dv.add(f'F2:F{ROWS + 1}')

info = wb.create_sheet('怎麼填')
info.column_dimensions['A'].width = 90
lines = [
    ('全班對戰・題庫範本', True),
    ('', False),
    ('1. 在「題目」分頁，一列填一題。前三列是範例，可以直接改掉或刪除。', False),
    ('2. 選項至少填 2 個（A、B），C、D 可以空白，例如是非題只填「⭕ 對」「❌ 錯」。', False),
    ('3. 正確答案填 A、B、C 或 D（也可以填 1～4，或直接寫正確選項的文字）。', False),
    ('4. 分數直接打 3/4 就好：這份範本的格子已經設成「文字」，不會被 Excel 變成日期。', False),
    ('   如果是自己新增的格子，請先選取格子 → 右鍵「儲存格格式」→ 選「文字」再打分數。', False),
    ('5. 題目要放圖片的話，先匯入，再到後台那一題按「🖼️ 加圖片」。', False),
    ('6. 存檔後，回到老師後台 →「🎮 題庫」→ 編輯題組 →「📋 整批貼上」→「📂 從檔案匯入」選這個檔案。', False),
    ('', False),
    ('小提醒：一個題組最多 100 題；第一列的標題「題目」會自動跳過。', False),
]
for i, (t, bold) in enumerate(lines, start=1):
    c = info.cell(row=i, column=1, value=t)
    c.font = Font(bold=bold, size=16 if bold else 12, color=ROSE if bold else INK)
    c.alignment = Alignment(vertical='center', wrap_text=True)

OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print(f'已產生 {OUT.relative_to(ROOT)}')
